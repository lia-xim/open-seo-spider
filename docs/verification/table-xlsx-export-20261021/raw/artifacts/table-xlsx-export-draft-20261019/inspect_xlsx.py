"""Independent read-only actual-byte oracle: stdlib ZIP/XML + installed openpyxl.

Does not import the JS writer. ST_Xstring is decoded once after XML decoding.
openpyxl's actual inline-string behavior is reported separately from OOXML.
"""
import hashlib
import json
import re
import sys
import warnings
import zipfile
import xml.etree.ElementTree as ET

import openpyxl

S = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
R = "{http://schemas.openxmlformats.org/package/2006/relationships}"
CT = "{http://schemas.openxmlformats.org/package/2006/content-types}"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
PARTS = ["[Content_Types].xml", "_rels/.rels", "xl/workbook.xml",
         "xl/_rels/workbook.xml.rels", "xl/worksheets/sheet1.xml"]


def col(n):
    text = ""
    while n:
        n, rest = divmod(n - 1, 26)
        text = chr(65 + rest) + text
    return text


def xstring(text):
    # A single scan, so protected literal patterns are never decoded recursively.
    return re.sub(r"_x([0-9A-Fa-f]{4})_", lambda m: chr(int(m[1], 16)), text)


def inspect(file, expected):
    with open(file, "rb") as stream:
        original = stream.read()
    with zipfile.ZipFile(file) as archive:
        assert archive.testzip() is None, "Bad ZIP CRC"
        assert archive.namelist() == PARTS, "Unexpected or duplicate package part"
        for info in archive.infolist():
            assert info.compress_type == zipfile.ZIP_STORED
            assert info.flag_bits & 0x800
            assert info.date_time == (1980, 1, 1, 0, 0, 0)
            assert not info.is_dir() and info.file_size == info.compress_size
        xml = {name: ET.fromstring(archive.read(name)) for name in PARTS}
    content_types = xml[PARTS[0]]
    assert content_types.tag == CT + "Types"
    assert {(n.attrib.get("PartName"), n.attrib.get("ContentType"))
            for n in content_types if n.tag == CT + "Override"} == {
        ("/xl/workbook.xml", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"),
        ("/xl/worksheets/sheet1.xml", "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml")}
    for name, target, kind in [("_rels/.rels", "xl/workbook.xml", "officeDocument"),
                               ("xl/_rels/workbook.xml.rels", "worksheets/sheet1.xml", "worksheet")]:
        rels = xml[name]
        assert rels.tag == R + "Relationships" and len(rels) == 1
        assert rels[0].attrib == {"Id": "rId1", "Type": REL + kind, "Target": target}
    sheet = xml["xl/worksheets/sheet1.xml"]
    assert sheet.tag == S + "worksheet"
    assert [c.tag for c in sheet] == [S + "dimension", S + "sheetData"]
    assert not sheet.findall(".//" + S + "f")
    assert not sheet.findall(".//" + S + "hyperlink")
    headers, records = expected["headers"], expected["rows"]
    assert len(headers) == len(set(headers)) and all(headers)
    matrix = [headers] + records
    last = col(len(headers)) + str(len(matrix))
    assert sheet[0].attrib == {"ref": "A1:" + last}
    assert len(sheet[1]) == len(matrix)
    xml_wire = {}
    cells = 0
    for index, (element, row) in enumerate(zip(sheet[1], matrix), 1):
        assert element.tag == S + "row" and element.attrib == {"r": str(index)}
        assert len(element) == len(headers) == len(row)
        for column, (cell, value) in enumerate(zip(element, row), 1):
            ref = col(column) + str(index)
            assert cell.tag == S + "c" and cell.attrib.get("r") == ref
            cells += 1
            if value is None:
                assert cell.attrib == {"r": ref} and len(cell) == 0
            elif isinstance(value, str):
                assert cell.attrib == {"r": ref, "t": "inlineStr"}
                assert len(cell) == 1 and cell[0].tag == S + "is"
                assert len(cell[0]) == 1 and cell[0][0].tag == S + "t"
                assert cell[0][0].attrib == {"{http://www.w3.org/XML/1998/namespace}space": "preserve"}
                wire = cell[0][0].text or ""
                assert xstring(wire) == value, (ref, repr(value), repr(wire))
                xml_wire[ref] = wire
            elif isinstance(value, bool):
                assert cell.attrib == {"r": ref, "t": "b"}
                assert len(cell) == 1 and cell[0].tag == S + "v"
                assert cell[0].text == ("1" if value else "0")
            else:
                assert isinstance(value, (int, float))
                assert cell.attrib == {"r": ref, "t": "n"}
                assert len(cell) == 1 and cell[0].tag == S + "v"
                assert float(cell[0].text) == value
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        workbook = openpyxl.load_workbook(file, data_only=False, read_only=False)
    assert workbook.sheetnames == ["Ergebnisse"]
    worksheet = workbook.active
    assert worksheet.max_row == len(matrix) and worksheet.max_column == len(headers)
    differences = []
    for index, row in enumerate(matrix, 1):
        for column, expected_value in enumerate(row, 1):
            cell = worksheet.cell(index, column)
            assert cell.hyperlink is None and cell.data_type != "f"
            if isinstance(expected_value, str):
                if expected_value == "" and cell.value is None:
                    differences.append({"cell": cell.coordinate, "kind": "literal_empty_collapsed_by_reader"})
                else:
                    assert cell.data_type == "s"
                    if cell.value != expected_value:
                        assert cell.value == xml_wire[cell.coordinate]
                        assert xstring(cell.value) == expected_value
                        differences.append({"cell": cell.coordinate, "kind": "inline_ST_Xstring_not_decoded_by_reader",
                                            "reader_value": cell.value, "specified_value": expected_value})
            elif expected_value is None:
                assert cell.value is None
            elif isinstance(expected_value, bool):
                assert type(cell.value) is bool and cell.value == expected_value
            else:
                assert type(cell.value) in (int, float) and cell.value == expected_value
    workbook.close()
    return {"result": "PASS", "kind": "independent_actual_ZIP_XML_and_openpyxl_read",
            "file": file, "bytes": len(original), "sha256": hashlib.sha256(original).hexdigest(),
            "parts": PARTS, "rows": len(records), "headers": headers, "decoded_cells": cells,
            "last_coordinate": last, "formula_elements": 0, "external_relationships": 0,
            "openpyxl": {"version": openpyxl.__version__, "warnings": [str(w.message) for w in caught],
                         "actual_reader_differences": differences},
            "native_Excel": "NOT_PROVEN", "scope": "Read-only format controls, not native/UI acceptance or schema-conformance certification."}


if __name__ == "__main__":
    if sys.argv[2] == "--reject":
        try:
            with zipfile.ZipFile(sys.argv[1]) as archive:
                assert archive.testzip() is None
        except (zipfile.BadZipFile, AssertionError):
            print(json.dumps({"result": "REJECTED_BAD_ZIP"}))
        else:
            raise AssertionError("Deliberately corrupted package was accepted")
    else:
        with open(sys.argv[2], encoding="utf8") as stream:
            expected = json.load(stream)
        print(json.dumps(inspect(sys.argv[1], expected), ensure_ascii=False))
